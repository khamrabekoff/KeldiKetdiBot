"""Advances: an employee records cash taken before payday, admins hear about
it, and the reports take it off the month's pay."""
import asyncio
import datetime
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# The Windows console is not UTF-8 by default and the screens carry emoji
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

os.environ['BOT_TOKEN'] = '123456:AAHtesttokenAAHtesttokenAAHtesttoken'
os.environ['ADMIN_SECRET'] = 'test_secret'
os.environ['DATABASE_PATH'] = os.path.join(tempfile.mkdtemp(), 'advances.db')

import openpyxl  # noqa: E402
from telegram.ext import ConversationHandler  # noqa: E402

import app  # noqa: E402
import database as db  # noqa: E402
import excel_export  # noqa: E402
import messages as msg  # noqa: E402
import settings  # noqa: E402
import ui  # noqa: E402
import utils  # noqa: E402

failures = []


def check(label, condition, detail=''):
    print(f"  {'PASS' if condition else 'FAIL'}  {label}{(' — ' + detail) if detail else ''}")
    if not condition:
        failures.append(label)


class FakeMessage:
    def __init__(self, text=''):
        self.text = text
        self.replies = []

    async def reply_text(self, text, **kwargs):
        self.replies.append(text)


class FakeUser:
    def __init__(self, user_id):
        self.id = user_id


class FakeUpdate:
    def __init__(self, user_id, text=''):
        self.message = FakeMessage(text)
        self.effective_user = FakeUser(user_id)


class FakeContext:
    def __init__(self):
        self.user_data = {}


class FakeBot:
    def __init__(self):
        self.sent = []

    async def send_message(self, chat_id, text, **kwargs):
        self.sent.append({'chat_id': chat_id, 'text': text, **kwargs})


class FakeApp:
    def __init__(self):
        self.bot = FakeBot()


class FakeQuery:
    def __init__(self, user_id, data):
        self.from_user = FakeUser(user_id)
        self.data = data
        self.answers = []
        self.edits = []

    async def answer(self, text=None):
        self.answers.append(text)

    async def edit_message_text(self, text, **kwargs):
        self.edits.append(text)


class FakeCallbackUpdate:
    def __init__(self, user_id, data):
        self.callback_query = FakeQuery(user_id, data)
        self.effective_user = FakeUser(user_id)


def say(handler, user_id, text, context):
    update = FakeUpdate(user_id, text)
    state = asyncio.run(handler(update, context))
    return state, update.message.replies


def press(user_id, data):
    update = FakeCallbackUpdate(user_id, data)
    asyncio.run(app.advance_void_callback(update, FakeContext()))
    return update.callback_query


EMP = 701
OTHER = 702
ADMIN = 703
ADMIN2 = 704

db.init_db()
db.add_user(EMP, '998900000020', 'Mustafo')
db.add_user(OTHER, '998900000021', 'Jasur')
db.add_user(ADMIN, '998900000022', 'Boss', role='admin')
db.add_user(ADMIN2, '998900000023', 'Boss 2', role='admin')
db.update_rates(EMP, 'per_minute', rate_per_minute=0.05)

now = utils.get_now()
month_start = now.replace(day=1).date()
# One finished day worth $27 (540 min x $0.05), so the balance is checkable.
day = now.replace(hour=9, minute=0, second=0, microsecond=0)
db.update_attendance_manual(EMP, day.date(), day, day.replace(hour=18), 27.0)

print("1. Кнопка и обработчики")
application = app.create_application()
keyboard_labels = [button.text for row in ui.employee_keyboard(EMP).keyboard for button in row]
check("кнопка «💸 Avans» у сотрудника", ui.BTN_ADVANCE in keyboard_labels, str(keyboard_labels))
admin_labels = [button.text for row in ui.admin_keyboard().keyboard for button in row]
check("у админа клавиатура не изменилась", ui.BTN_ADVANCE not in admin_labels)

names = [getattr(h, 'name', None) for group in application.handlers.values() for h in group]
check("диалог аванса зарегистрирован", 'advance_conv' in names)
patterns = [getattr(h, 'pattern', None) for group in application.handlers.values() for h in group]
patterns = [p.pattern for p in patterns if p is not None]
check("кнопка отмены у админа зарегистрирована", '^adv:void:' in patterns)
positions = {getattr(h, 'name', None) or h.callback.__name__: i
             for i, h in enumerate(application.handlers[0])}
check("диалог стоит раньше общего текстового обработчика",
      positions['advance_conv'] < positions['unknown_text'])

# From here on messages go to a recorder instead of Telegram
app.telegram_app = FakeApp()
sent = app.telegram_app.bot.sent

print("\n2. Сотрудник записывает аванс с описанием")
ctx = FakeContext()
state, replies = say(app.advance_start, EMP, ui.BTN_ADVANCE, ctx)
check("спрашивает сумму", state == app.ADV_AMOUNT and 'Qancha' in replies[0])

state, replies = say(app.advance_amount, EMP, 'uch yuz', ctx)
check("не число — просит ещё раз", state == app.ADV_AMOUNT and 'raqam' in replies[0])

state, replies = say(app.advance_amount, EMP, '$300', ctx)
check("сумма принята, спрашивает описание", state == app.ADV_NOTE and 'Izoh' in replies[0])

state, replies = say(app.advance_note, EMP, 'Ijara <tez>', ctx)
check("показывает подтверждение", state == app.ADV_CONFIRM and '$300.00' in replies[0])
check("описание экранировано", '&lt;tez&gt;' in replies[0])

state, replies = say(app.advance_confirm, EMP, ui.BTN_ADVANCE_CONFIRM, ctx)
check("диалог завершён", state == ConversationHandler.END)
check("сотруднику подтверждение", 'Avans qayd etildi: $300.00' in replies[0])

rows = db.get_advances(month_start, utils.next_month(month_start), EMP)
check("запись в базе одна", len(rows) == 1)
check("сумма и описание сохранены", rows and rows[0]['amount'] == 300 and rows[0]['note'] == 'Ijara <tez>')
first_id = rows[0]['id'] if rows else None

print("\n3. Уведомление админам")
to_admins = [m for m in sent if m['chat_id'] in (ADMIN, ADMIN2)]
check("пришло обоим админам", len(to_admins) == 2)
check("сотруднику-коллеге ничего", not [m for m in sent if m['chat_id'] == OTHER])
card = to_admins[0]['text'] if to_admins else ''
check("в карточке имя, сумма, описание",
      'Mustafo' in card and '$300.00' in card and 'Ijara &lt;tez&gt;' in card, card.replace('\n', ' | '))
button = to_admins[0]['reply_markup'].inline_keyboard[0][0] if to_admins else None
check("кнопка отмены ведёт на этот аванс", button and button.callback_data == f"adv:void:{first_id}")
print("     " + card.replace('\n', '\n     '))

print("\n4. Описание не обязательно")
sent.clear()
ctx = FakeContext()
say(app.advance_start, EMP, ui.BTN_ADVANCE, ctx)
say(app.advance_amount, EMP, '100', ctx)
state, replies = say(app.advance_note, EMP, ui.BTN_ADVANCE_SKIP_NOTE, ctx)
check("без описания в подтверждении нет строки Izoh", 'Izoh' not in replies[0])
say(app.advance_confirm, EMP, ui.BTN_ADVANCE_CONFIRM, ctx)
rows = db.get_advances(month_start, utils.next_month(month_start), EMP)
check("вторая запись без описания", len(rows) == 2 and rows[1]['note'] is None)
check("у админа итог за месяц $400.00", sent and '$400.00' in sent[0]['text'])

print("\n5. «Ortga» отменяет на любом шаге")
for step_text in ([], ['50'], ['50', 'izoh']):
    ctx = FakeContext()
    say(app.advance_start, EMP, ui.BTN_ADVANCE, ctx)
    handlers = [app.advance_amount, app.advance_note, app.advance_confirm]
    for handler, text in zip(handlers, step_text):
        say(handler, EMP, text, ctx)
    state, replies = say(handlers[len(step_text)], EMP, msg.BTN_BACK, ctx)
    check(f"отмена после {len(step_text)} шагов", state == ConversationHandler.END and 'Bekor' in replies[0])
check("лишних записей нет", len(db.get_advances(month_start, utils.next_month(month_start), EMP)) == 2)

state, replies = say(app.advance_start, ADMIN, ui.BTN_ADVANCE, FakeContext())
check("админ диалог не начинает", state == ConversationHandler.END and not replies)

print("\n6. Где видно")
stats = ui.employee_stats_card(EMP)
check("Hisobim: сумма авансов", 'Avans: <b>$400.00</b>' in stats)
check("Hisobim: остаток 27 − 400", 'Qoldiq: <b>-$373.00</b>' in stats, 'отрицательный остаток')
check("Hisobim: список с описанием", 'Ijara &lt;tez&gt;' in stats)

settings.set_bool('show_employee_earnings', False)
stats = ui.employee_stats_card(EMP)
check("зарплата скрыта: аванс виден", 'Avans: <b>$400.00</b>' in stats)
check("зарплата скрыта: остатка нет", 'Qoldiq' not in stats)
settings.set_bool('show_employee_earnings', True)

emp_card = ui.admin_employee_card(EMP)
check("карточка у админа: аванс и остаток", 'Avans:   $400.00' in emp_card and 'Qoldiq:  -$373.00' in emp_card)

report = ui.admin_month_report_card(month_start)
check("месячный отчёт: строка аванса", '− 400.00 avans' in report)
check("месячный отчёт: сколько заплатить сотруднику", "= -$373.00 to'lash kerak" in report)
check("месячный отчёт: итоги", 'Avans: $400.00' in report and "TO'LASH KERAK: -$373.00" in report)
print("     " + report.replace('\n', '\n     '))

other_card = ui.admin_employee_card(OTHER)
check("у сотрудника без авансов карточка без них", 'Avans' not in other_card)

print("\n7. Прошлый месяц не смешивается")
last_month_day = datetime.datetime.combine(month_start - datetime.timedelta(days=3), datetime.time(12, 0))
db.add_advance(EMP, 50, 'eski', last_month_day)
check("в этом месяце всё ещё $400", 'Avans: <b>$400.00</b>' in ui.employee_stats_card(EMP))
last_report = ui.admin_month_report_card(last_month_day.date().replace(day=1))
check("в отчёте прошлого месяца свой аванс $50", 'Avans: $50.00' in last_report)
check("и нет дней этого месяца", 'JAMI: $0.00' in last_report and '0 kun' in last_report,
      'раньше отчёт за прошлый месяц считал и текущий')
check("к выплате за прошлый месяц -$50", "TO'LASH KERAK: -$50.00" in last_report)

print("\n8. Админ отменяет аванс")
sent.clear()
query = press(OTHER, f"adv:void:{first_id}")
check("сотрудник не может отменить", db.get_advance(first_id)['status'] == 'ACTIVE' and not query.edits)

query = press(ADMIN, f"adv:void:{first_id}")
check("отменён", db.get_advance(first_id)['status'] == 'VOID')
check("кто отменил, записано", db.get_advance(first_id)['voided_by'] == ADMIN)
check("сообщение у админа помечено", query.edits and 'Bekor qilingan' in query.edits[0])
check("сотрудник узнал", [m for m in sent if m['chat_id'] == EMP and '$300.00' in m['text']])

sent.clear()
query = press(ADMIN2, f"adv:void:{first_id}")
check("второй админ: уже отменён", query.answers == ['Allaqachon bekor qilingan'])
check("сотруднику второй раз не пишем", not sent)

check("итог за месяц уменьшился до $100", 'Avans: <b>$100.00</b>' in ui.employee_stats_card(EMP))

print("\n9. Excel")
db.add_advance(OTHER, 20, 'Проезд', now)  # Cyrillic note, no days worked
bio = excel_export.create_monthly_report_excel(month_start)
check("файл создан", bio is not None)
wb = openpyxl.load_workbook(bio)
check("первый лист — «To'lov»", wb.sheetnames[0] == "To'lov" and wb.active.title == "To'lov", str(wb.sheetnames))
pay = wb["To'lov"]
table = {pay.cell(row=r, column=1).value: [pay.cell(row=r, column=c).value for c in range(2, 6)]
         for r in range(4, pay.max_row + 1) if pay.cell(row=r, column=1).value}
check("Mustafo: дни, заработок, аванс, к выплате", table.get('Mustafo') == [1, 27.0, 100.0, -73.0],
      str(table.get('Mustafo')))
check("Jasur без дней, только аванс", table.get('Jasur') == [0, 0.0, 20.0, -20.0], str(table.get('Jasur')))
check("JAMI к выплате", table.get('JAMI') and table['JAMI'][1:] == [27.0, 120.0, -93.0], str(table.get('JAMI')))
check("отменённый аванс не попал", 300.0 not in [v for row in table.values() for v in row])

check("лист «Avanslar» есть", 'Avanslar' in wb.sheetnames)
adv_sheet = wb['Avanslar']
notes = [adv_sheet.cell(row=r, column=5).value for r in range(4, adv_sheet.max_row + 1)]
check("описания в листе авансов", 'Проезд' in notes, str(notes))

main = wb['Oylik Hisobot']
main_values = [c.value for row in main.iter_rows() for c in row]
check("в основном листе итог с колонкой «To'lash kerak»", any("To'lash kerak" in str(v) for v in main_values))

print("\n10. PDF")
import pdf_export  # noqa: E402
pdf = pdf_export.create_monthly_pdf_report(month_start)
check("PDF создан", pdf is not None and pdf.getvalue()[:4] == b'%PDF')
check("шрифт с кириллицей подключён", pdf_export._unicode_fonts()[0] == 'KKSans', str(pdf_export._unicode_fonts()))
pdf_last = pdf_export.create_monthly_pdf_report(last_month_day.date().replace(day=1))
check("PDF за месяц только с авансом тоже создаётся", pdf_last is not None)

print("\n11. Разбор суммы и вывод денег")
cases = {
    '300': 300.0, '$300': 300.0, '300 $': 300.0, '1 500': 1500.0, '1,500': 1500.0,
    '12.5': 12.5, '12,5': 12.5, '0': None, '-5': None, 'abc': None, '': None, 'nan': None,
}
for text, expected in cases.items():
    got = utils.parse_amount(text)
    check(f"parse_amount({text!r}) = {expected}", got == expected, f"получено {got}")
check("минус перед знаком валюты", utils.format_money(-50) == '-$50.00', utils.format_money(-50))
check("положительные суммы как раньше", utils.format_money(208333.333) == '$208,333.33')

print("\n" + ("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ" if not failures else f"ПАДЕНИЙ: {len(failures)} -> {failures}"))
sys.exit(0 if not failures else 1)

"""Salary payments and advances entered by an admin.

The admin marks pay as paid from the monthly report; a settled month hands
whatever is left of its balance to the next one - and only a settled month,
or every month paid outside the bot would show up as a debt.
"""
import asyncio
import datetime
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# The Windows console is not UTF-8 by default and the screens carry emoji
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

os.environ['BOT_TOKEN'] = '123456:AAHtesttokenAAHtesttokenAAHtesttoken'
os.environ['ADMIN_SECRET'] = 'test_secret'
os.environ['DATABASE_PATH'] = os.path.join(tempfile.mkdtemp(), 'payments.db')

import openpyxl  # noqa: E402
from telegram.ext import ConversationHandler  # noqa: E402

import analytics  # noqa: E402
import app  # noqa: E402
import database as db  # noqa: E402
import excel_export  # noqa: E402
import pdf_export  # noqa: E402
import ui  # noqa: E402
import utils  # noqa: E402

failures = []


def check(label, condition, detail=''):
    print(f"  {'PASS' if condition else 'FAIL'}  {label}{(' — ' + detail) if detail else ''}")
    if not condition:
        failures.append(label)


class FakeMessage:
    def __init__(self, text=''):
        self.text = text
        self.replies = []

    async def reply_text(self, text, **kwargs):
        self.replies.append({'text': text, **kwargs})


class FakeUser:
    def __init__(self, user_id):
        self.id = user_id


class FakeUpdate:
    def __init__(self, user_id, text=''):
        self.message = FakeMessage(text)
        self.effective_user = FakeUser(user_id)


class FakeContext:
    def __init__(self):
        self.user_data = {}


class FakeBot:
    def __init__(self):
        self.sent = []

    async def send_message(self, chat_id, text, **kwargs):
        self.sent.append({'chat_id': chat_id, 'text': text, **kwargs})


class FakeApp:
    def __init__(self):
        self.bot = FakeBot()


class FakeQuery:
    def __init__(self, user_id, data):
        self.from_user = FakeUser(user_id)
        self.data = data
        self.answers = []
        self.edits = []
        self.message = FakeMessage()

    async def answer(self, text=None):
        self.answers.append(text)

    async def edit_message_text(self, text, **kwargs):
        self.edits.append({'text': text, **kwargs})


class FakeCallbackUpdate:
    def __init__(self, user_id, data):
        self.callback_query = FakeQuery(user_id, data)
        self.effective_user = FakeUser(user_id)


def press(handler, user_id, data, context=None):
    update = FakeCallbackUpdate(user_id, data)
    result = asyncio.run(handler(update, context or FakeContext()))
    return update.callback_query, result


def say(handler, user_id, text, context):
    update = FakeUpdate(user_id, text)
    state = asyncio.run(handler(update, context))
    return state, update.message.replies


def callbacks(markup):
    return [b.callback_data for row in markup.inline_keyboard for b in row]


def button_for(markup, prefix):
    for row in markup.inline_keyboard:
        for b in row:
            if b.callback_data.startswith(prefix):
                return b
    return None


EMP = 801
EMP2 = 802
ADMIN = 803
ADMIN2 = 804

db.init_db()
db.add_user(EMP, '998900000030', 'Ibrohim')
db.add_user(EMP2, '998900000031', 'Mohir')
db.add_user(ADMIN, '998900000032', 'Boss', role='admin')
db.add_user(ADMIN2, '998900000033', 'Boss 2', role='admin')
db.update_rates(EMP, 'per_minute', rate_per_minute=0.05)
db.update_rates(EMP2, 'per_minute', rate_per_minute=0.05)

now = utils.get_now()
THIS = now.replace(day=1).date()
PREV = (THIS - datetime.timedelta(days=1)).replace(day=1)
OLDER = (PREV - datetime.timedelta(days=1)).replace(day=1)


def work_day(user_id, day, wage):
    start = datetime.datetime.combine(day, datetime.time(9, 0))
    db.update_attendance_manual(user_id, day, start, start.replace(hour=18), wage)


# OLDER: paid outside the bot, never settled here. PREV: $27 earned, $100 advance.
work_day(EMP, OLDER.replace(day=5), 500.0)
work_day(EMP, PREV.replace(day=5), 27.0)
db.add_advance(EMP, 100, 'eski avans', datetime.datetime.combine(PREV.replace(day=6), datetime.time(12)))
work_day(EMP, THIS, 200.0)
work_day(EMP2, THIS, 50.0)

app.create_application()
app.telegram_app = FakeApp()
sent = app.telegram_app.bot.sent

print("1. Обработчики и кнопки")
application = app.create_application()
app.telegram_app = FakeApp()
sent = app.telegram_app.bot.sent
patterns = []
for group in application.handlers.values():
    for handler in group:
        for h in [handler] + list(getattr(handler, 'entry_points', [])):
            pattern = getattr(h, 'pattern', None)
            if pattern is not None:
                patterns.append(pattern.pattern)
for needed in ['^pay:list:', '^pay:emp:', '^pay:do:', '^pay:void:', '^pay:sum:', '^advg:']:
    check(f"зарегистрирован {needed}", needed in patterns)

ym_this, ym_prev = THIS.strftime('%Y-%m'), PREV.strftime('%Y-%m')
check("в отчёте есть кнопка «💰 To'lov»", f"pay:list:{ym_this}" in callbacks(ui.admin_report_keyboard(ym_this)))
query, _ = press(app.handle_callback_query, ADMIN, f"edit_{EMP}")
check("в карточке сотрудника есть «💸 Avans berish»",
      query.edits and f"advg:{EMP}" in callbacks(query.edits[0]['reply_markup']))

print("\n2. Месяц, оплаченный мимо бота, ничего не переносит")
older = analytics.employee_month(EMP, OLDER)
check("OLDER не закрыт", older and not older['settled'])
prev = analytics.employee_month(EMP, PREV)
check("в PREV нет переноса из OLDER", prev['carry_in'] == 0, f"carry_in={prev['carry_in']}")
check("PREV к выплате 27 − 100 = −73", round(prev['to_pay'], 2) == -73.0)

print("\n3. Закрываем PREV с долгом")
text, markup = ui.payment_employee_card(EMP, PREV)
close_button = button_for(markup, 'pay:do:')
check("кнопка «Oyni yopish»", close_button and 'yopish' in close_button.text)
check("объяснение про долг", 'keyingi oy' in text)

query, _ = press(app.payment_do_callback, EMP2, close_button.callback_data)
check("сотрудник не может закрыть месяц", not db.get_payments(PREV))

query, _ = press(app.payment_do_callback, ADMIN, close_button.callback_data)
check("месяц закрыт записью на $0", [p['amount'] for p in db.get_payments(PREV, EMP)] == [0])
check("сотруднику сообщили о долге", any(m['chat_id'] == EMP and '73.00' in m['text'] for m in sent))
check("экран перерисован", query.edits and 'Oy yopilgan' in query.edits[-1]['text'])

query, _ = press(app.payment_do_callback, ADMIN2, close_button.callback_data)
check("второй раз не закрывается", len(db.get_payments(PREV, EMP)) == 1 and 'o\'zgardi' in (query.answers[-1] or ''))

print("\n4. Долг переходит в текущий месяц")
this = analytics.employee_month(EMP, THIS)
check("перенос −73", round(this['carry_in'], 2) == -73.0, f"{this['carry_in']}")
check("к выплате 200 − 73 = 127", round(this['to_pay'], 2) == 127.0, f"{this['to_pay']}")
report = ui.admin_month_report_card(THIS)
check("в отчёте строка переноса", "− 73.00 o'tgan oydan" in report)
check("в отчёте итог к выплате 127 + 50", "TO'LASH KERAK: $177.00" in report, report.split('━')[-1])
stats = ui.employee_stats_card(EMP)
check("сотрудник видит перенос и остаток", "O'tgan oydan: <b>-$73.00</b>" in stats and 'Qoldiq: <b>$127.00</b>' in stats)

print("\n5. Выплата полной суммы")
sent.clear()
text, markup = ui.payment_employee_card(EMP, THIS)
pay_button = button_for(markup, 'pay:do:')
check("кнопка «$127.00 to'landi»", pay_button and '$127.00' in pay_button.text)
stale = pay_button.callback_data.rsplit(':', 1)[0] + ':12600'
query, _ = press(app.payment_do_callback, ADMIN, stale)
check("устаревшая сумма не записывается", not db.get_payments(THIS))
query, _ = press(app.payment_do_callback, ADMIN, pay_button.callback_data)
check("записано $127", [p['amount'] for p in db.get_payments(THIS, EMP)] == [127.0])
check("сотрудник получил сообщение", any(m['chat_id'] == EMP and '$127.00' in m['text'] for m in sent))
check("карточка: полностью оплачено", query.edits and "To'liq to'langan" in query.edits[-1]['text'])
query, _ = press(app.payment_do_callback, ADMIN2, pay_button.callback_data)
check("двойная выплата невозможна", len(db.get_payments(THIS, EMP)) == 1)

list_text, list_markup = ui.payment_list_card(THIS)
labels = [b.text for row in list_markup.inline_keyboard for b in row]
check("в списке Ibrohim отмечен ✅", any(l.startswith('✅ Ibrohim') for l in labels), str(labels))
check("Mohir ещё к оплате $50", any('Mohir · $50.00' in l for l in labels))
report = ui.admin_month_report_card(THIS)
check("в отчёте «to'langan» и ✅", "− 127.00 to'langan" in report and "to'lash kerak</code> ✅" in report)

print("\n6. Другая сумма (частичная выплата)")
ctx = FakeContext()
query, state = press(app.payment_sum_start, ADMIN, f"pay:sum:{EMP2}:{ym_this}", ctx)
check("спрашивает сумму", state == app.PAY_SUM and query.message.replies)
state, replies = say(app.payment_sum_value, ADMIN, 'ellik', ctx)
check("не число — просит ещё раз", state == app.PAY_SUM)
state, replies = say(app.payment_sum_value, ADMIN, '30', ctx)
check("записано $30", state == ConversationHandler.END and
      [p['amount'] for p in db.get_payments(THIS, EMP2)] == [30.0])
check("осталось $20", round(analytics.employee_month(EMP2, THIS)['to_pay'], 2) == 20.0)
check("админу вернулась клавиатура", replies and replies[0].get('reply_markup') is not None)

print("\n7. Отмена выплаты")
sent.clear()
payment_id = db.get_payments(THIS, EMP2)[0]['id']
query, _ = press(app.payment_void_callback, EMP, f"pay:void:{payment_id}")
check("сотрудник не может отменить", db.get_payments(THIS, EMP2))
query, _ = press(app.payment_void_callback, ADMIN, f"pay:void:{payment_id}")
check("выплата отменена", not db.get_payments(THIS, EMP2))
check("к выплате снова $50", round(analytics.employee_month(EMP2, THIS)['to_pay'], 2) == 50.0)
check("сотруднику сообщили", any(m['chat_id'] == EMP2 for m in sent))

print("\n8. Excel и PDF с выплатами")
bio = excel_export.create_monthly_report_excel(THIS)
pay = openpyxl.load_workbook(bio)["To'lov"]
header = [pay.cell(row=3, column=c).value for c in range(1, pay.max_column + 1)]
check("колонки переноса и выплат появились",
      any("O'tgan oydan" in str(h) for h in header) and any("To'langan" in str(h) for h in header), str(header))
rows = {pay.cell(row=r, column=1).value: [pay.cell(row=r, column=c).value for c in range(2, pay.max_column + 1)]
        for r in range(4, pay.max_row + 1) if pay.cell(row=r, column=1).value}
check("Ibrohim: 200, −73, 0, 127, 0", rows.get('Ibrohim') == [1, 200, -73, 0, 127, 0], str(rows.get('Ibrohim')))
formula = pay.cell(row=pay.max_row, column=1).value
check("под таблицей формула с переносом", "O'tgan oydan" in str(formula), str(formula))
check("PDF создаётся", pdf_export.create_monthly_pdf_report(THIS) is not None)

print("\n9. Аванс, который записывает админ")
sent.clear()
ctx = FakeContext()
query, state = press(app.advance_admin_start, ADMIN, f"advg:{EMP2}", ctx)
check("спрашивает сумму для Mohir", state == app.ADV_AMOUNT and 'Mohir' in query.message.replies[0]['text'])
say(app.advance_amount, ADMIN, '40', ctx)
state, replies = say(app.advance_note, ADMIN, "Yo'l kira", ctx)
check("в подтверждении имя сотрудника", 'Mohir' in replies[0]['text'])
state, replies = say(app.advance_confirm, ADMIN, ui.BTN_ADVANCE_CONFIRM, ctx)
check("диалог завершён", state == ConversationHandler.END)
advances = db.get_advances(THIS, utils.next_month(THIS), EMP2)
check("аванс записан на Mohir", len(advances) == 1 and advances[0]['amount'] == 40)
check("записано, кто внёс", db.get_advance(advances[0]['id'])['created_by'] == ADMIN)
check("Mohir получил сообщение", any(m['chat_id'] == EMP2 and 'Sizga avans yozildi' in m['text'] for m in sent))
to_admin2 = [m for m in sent if m['chat_id'] == ADMIN2]
check("второй админ получил карточку с автором", to_admin2 and 'Admin yozdi' in to_admin2[0]['text'])
check("автору карточка не дублируется", not [m for m in sent if m['chat_id'] == ADMIN])
check("автору подтверждение и его клавиатура",
      'Avans yozildi' in replies[0]['text'] and replies[0].get('reply_markup') is not None)
check("к выплате Mohir 50 − 40 = 10", round(analytics.employee_month(EMP2, THIS)['to_pay'], 2) == 10.0)

ctx = FakeContext()
query, state = press(app.advance_admin_start, EMP, f"advg:{EMP2}", ctx)
check("сотрудник не может записать аванс другому", state == ConversationHandler.END)

ctx = FakeContext()
press(app.advance_admin_start, ADMIN, f"advg:{EMP2}", ctx)
state, replies = say(app.advance_amount, ADMIN, 'Ortga', ctx)
check("«Ortga» у админа возвращает его меню",
      state == ConversationHandler.END and replies[0].get('reply_markup') is not None)

print("\n" + ("ВСЕ ПРОВЕРКИ ПРОЙДЕНЫ" if not failures else f"ПАДЕНИЙ: {len(failures)} -> {failures}"))
sys.exit(0 if not failures else 1)
